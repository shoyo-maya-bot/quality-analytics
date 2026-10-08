"""サンプルの入力（架空の会員・注文システムのレビュー指摘）を生成する.

    python samples/make_samples.py

- samples/指摘一覧.xlsx       指摘表（要件定義・基本設計のレビュー。列は config/sources/table.yaml に対応）
- samples/github_pulls.json   GitHub の PR レビューコメント（collect github --file で読める形）
- samples/office-review/      office-review のレビュー結果と判定（同リポジトリのサンプルと同じ形式）

人名・社名は使わず「レビュアーA」などの記号にしている。
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent

# (成果物, 工程, 指摘者, 作成者, 重要度, 指摘内容, 判定)
ROWS = [
    (
        "要件定義書_会員管理",
        "要件定義",
        "レビュアーA",
        "作成者1",
        "必須",
        "退会した会員のデータをいつまで保持するかの記載が無い",
        "対応する",
    ),
    (
        "要件定義書_会員管理",
        "要件定義",
        "レビュアーA",
        "作成者1",
        "必須",
        "会員登録の業務フローで、本人確認の手順が書かれていない",
        "対応する",
    ),
    ("要件定義書_会員管理", "要件定義", "レビュアーB", "作成者1", "推奨", "「適宜」通知するとあり、通知の条件が決まらない", "対応する"),
    ("要件定義書_会員管理", "要件定義", "レビュアーB", "作成者1", "軽微", "「メールアドレス」と「メアド」の表記ゆれ", "対応する"),
    ("要件定義書_会員管理", "要件定義", "レビュアーC", "作成者1", "質問", "会員数はどの程度を想定していますか？", "対応する"),
    ("要件定義書_会員管理", "要件定義", "レビュアーC", "作成者1", "推奨", "画面の項目は左右 2 列にした方が見やすい", "対応しない"),
    (
        "要件定義書_会員管理",
        "要件定義",
        "レビュアーA",
        "作成者1",
        "必須",
        "性能要件の応答時間が上位の企画書（3 秒）と異なる（5 秒）",
        "対応する",
    ),
    ("要件定義書_会員管理", "要件定義", "レビュアーB", "作成者1", "軽微", "見出しのフォントがそろっていない", "対応する"),
    ("要件定義書_注文管理", "要件定義", "レビュアーA", "作成者2", "必須", "注文のキャンセル期限の前提が書かれていない", "対応する"),
    (
        "要件定義書_注文管理",
        "要件定義",
        "レビュアーA",
        "作成者2",
        "必須",
        "外部の決済システムとの連携のタイミングが定義されていない",
        "対応する",
    ),
    (
        "要件定義書_注文管理",
        "要件定義",
        "レビュアーB",
        "作成者2",
        "推奨",
        "用語「確定」の定義が曖昧で、決済完了か出荷指示かで解釈が分かれる",
        "対応する",
    ),
    ("要件定義書_注文管理", "要件定義", "レビュアーC", "作成者2", "推奨", "業務フロー図は横向きにした方が読みやすくなる", "対応しない"),
    ("要件定義書_注文管理", "要件定義", "レビュアーC", "作成者2", "軽微", "句読点が「，」と「、」で混在", "対応する"),
    (
        "要件定義書_注文管理",
        "要件定義",
        "レビュアーB",
        "作成者2",
        "必須",
        "移行時に旧システムの注文データをどう扱うかの記載が無い",
        "対応する",
    ),
    ("要件定義書_注文管理", "要件定義", "レビュアーA", "作成者2", "質問", "1 日の注文件数の上限はどの程度ですか？", "対応する"),
    (
        "基本設計書_会員API",
        "基本設計",
        "レビュアーA",
        "作成者3",
        "必須",
        "会員照会 API で、存在しない会員 ID のときのステータスとメッセージ ID が無い",
        "対応する",
    ),
    (
        "基本設計書_会員API",
        "基本設計",
        "レビュアーA",
        "作成者3",
        "必須",
        "memberId を指定すれば他人の情報が取得できる。認可の記載が無い",
        "対応する",
    ),
    (
        "基本設計書_会員API",
        "基本設計",
        "レビュアーB",
        "作成者3",
        "必須",
        "氏名の桁が API 定義書（50）と画面設計書（40）で一致しない",
        "対応する",
    ),
    ("基本設計書_会員API", "基本設計", "レビュアーB", "作成者3", "軽微", "項目名の全角・半角の表記がそろっていない", "対応する"),
    ("基本設計書_会員API", "基本設計", "レビュアーC", "作成者3", "推奨", "レスポンスの項目順は ID を先頭にした方が好み", "対応しない"),
    ("基本設計書_会員API", "基本設計", "レビュアーC", "作成者3", "軽微", "誤字: 「照会」が「紹介」になっている", "対応する"),
    ("基本設計書_会員API", "基本設計", "レビュアーC", "作成者3", "推奨", "エラーメッセージの文言は既に記載済みのため対象外", "対応しない"),
    (
        "基本設計書_会員API",
        "基本設計",
        "レビュアーA",
        "作成者3",
        "必須",
        "メッセージ ID の採番が命名規則（APPE + 4 桁）に従っていない",
        "対応する",
    ),
    (
        "基本設計書_注文API",
        "基本設計",
        "レビュアーA",
        "作成者3",
        "必須",
        "注文検索で「件数が多い場合は適宜ページングする」とあり、1 ページの件数が決まらない",
        "対応する",
    ),
    (
        "基本設計書_注文API",
        "基本設計",
        "レビュアーA",
        "作成者3",
        "必須",
        "検索条件の日付の境界（タイムゾーン）の前提が書かれていない",
        "対応する",
    ),
    (
        "基本設計書_注文API",
        "基本設計",
        "レビュアーB",
        "作成者3",
        "必須",
        "注文登録 API のトランザクションの範囲が書かれていない",
        "対応する",
    ),
    ("基本設計書_注文API", "基本設計", "レビュアーB", "作成者3", "推奨", "参照テーブルと使うインデックスの記載が不足", "対応する"),
    ("基本設計書_注文API", "基本設計", "レビュアーC", "作成者3", "軽微", "表の罫線の太さがそろっていない（体裁）", "対応する"),
    ("基本設計書_注文API", "基本設計", "レビュアーC", "作成者3", "推奨", "API 名は英語にした方がよいと思う（提案）", "対応しない"),
    ("基本設計書_注文API", "基本設計", "レビュアーC", "作成者3", "軽微", "インデントがずれている", "対応する"),
    (
        "基本設計書_注文API",
        "基本設計",
        "レビュアーC",
        "作成者3",
        "質問",
        "同時に同じ注文を更新した場合の扱いはどうなりますか？",
        "対応する",
    ),
    (
        "基本設計書_注文取込バッチ",
        "基本設計",
        "レビュアーA",
        "作成者4",
        "必須",
        "取込ファイルが 0 件のときの終了コードが定義されていない",
        "対応する",
    ),
    (
        "基本設計書_注文取込バッチ",
        "基本設計",
        "レビュアーA",
        "作成者4",
        "必須",
        "再実行時に同じ注文番号を二重登録しない前提が書かれていない",
        "対応する",
    ),
    (
        "基本設計書_注文取込バッチ",
        "基本設計",
        "レビュアーA",
        "作成者4",
        "必須",
        "ジョブの処理ステップの順序が処理フロー図と異なる",
        "対応する",
    ),
    ("基本設計書_注文取込バッチ", "基本設計", "レビュアーB", "作成者4", "推奨", "異常終了時の監視とアラートの記載が無い", "対応する"),
    ("基本設計書_注文取込バッチ", "基本設計", "レビュアーB", "作成者4", "軽微", "「取込」と「取り込み」の表記ゆれ", "対応する"),
    ("基本設計書_注文取込バッチ", "基本設計", "レビュアーC", "作成者4", "推奨", "ログの出力形式はJSONの方が好み", "対応しない"),
    (
        "基本設計書_注文取込バッチ",
        "基本設計",
        "レビュアーA",
        "作成者4",
        "必須",
        "取込件数の上限と、超えたときの扱いが書かれていない",
        "対応する",
    ),
]

PULLS = {
    "repo": "example-org/member-service",
    "pulls": [
        {
            "number": 12,
            "title": "会員照会 API を追加",
            "user": {"login": "dev-1", "type": "User"},
            "comments": [
                {
                    "id": 1001,
                    "user": {"login": "reviewer-x", "type": "User"},
                    "path": "src/main/java/com/example/app/member/MemberService.java",
                    "body": "[MUST] 会員が存在しないときに null を返していて、設計書の APPE0001 の例外になっていない",
                    "created_at": "2026-09-20T10:00:00Z",
                },
                {
                    "id": 1002,
                    "user": {"login": "dev-1", "type": "User"},
                    "in_reply_to_id": 1001,
                    "body": "NotFoundException に直しました",
                    "created_at": "2026-09-20T11:00:00Z",
                },
                {
                    "id": 1003,
                    "user": {"login": "reviewer-x", "type": "User"},
                    "path": "src/test/java/com/example/app/member/MemberServiceTest.java",
                    "body": "[WANT] 退会会員のケースのテストが無い",
                    "created_at": "2026-09-20T10:05:00Z",
                },
                {
                    "id": 1004,
                    "user": {"login": "reviewer-y", "type": "User"},
                    "path": "src/main/java/com/example/app/member/MemberController.java",
                    "body": "[NITS] 変数名 m は member の方が読みやすい",
                    "created_at": "2026-09-20T10:10:00Z",
                },
                {
                    "id": 1005,
                    "user": {"login": "ci-bot", "type": "Bot"},
                    "path": "build.gradle",
                    "body": "Coverage 82%",
                    "created_at": "2026-09-20T10:11:00Z",
                },
            ],
            "reviews": [
                {
                    "id": 2001,
                    "user": {"login": "reviewer-y", "type": "User"},
                    "body": "[Q] 会員 ID の桁チェックはコントローラとサービスのどちらで行う方針ですか？",
                    "submitted_at": "2026-09-20T10:20:00Z",
                },
                {"id": 2002, "user": {"login": "reviewer-x", "type": "User"}, "body": "", "submitted_at": "2026-09-20T12:00:00Z"},
            ],
        },
        {
            "number": 15,
            "title": "注文取込ジョブの再実行対応",
            "user": {"login": "dev-2", "type": "User"},
            "comments": [
                {
                    "id": 1101,
                    "user": {"login": "reviewer-x", "type": "User"},
                    "path": "src/main/resources/mapper/OrderMapper.xml",
                    "body": "[MUST] SQL に ${} を使っていて規約違反（#{} にする）",
                    "created_at": "2026-09-25T09:00:00Z",
                },
                {
                    "id": 1102,
                    "user": {"login": "reviewer-y", "type": "User"},
                    "path": "src/main/java/com/example/app/order/batch/OrderImportJob.java",
                    "body": "ここは early return にした方がよさそう",
                    "created_at": "2026-09-25T09:30:00Z",
                },
            ],
            "reviews": [],
        },
    ],
}

REVIEW_MD = """---
target: inputs/api/order-api.md
type: api
reviewed_at: 2026-10-02
reviewer: "Copilot（確認: <担当者>）"
---

# レビュー結果: order-api

| No | check_id | 重要度 | ロール | 箇所 | 指摘 | 根拠 | 無効となるケースの確認 |
|---|---|---|---|---|---|---|---|
| 1 | ALL-003 | MUST | spec_manager | 3. 処理 3 | 「適宜ページングする」では 1 ページの件数と並び順の安定性が決まらない | 3. 処理 | 仕様の記載のため該当しない |
| 2 | API-003 | MUST | spec_manager | 4. 出力 | 該当なし・入力誤り時のステータスとメッセージ ID が無い | 4. 出力 | 参照先が無い |
| 3 | API-006 | MUST | infra | 2. 入力 | memberId を指定すれば他人の注文も取得できる。認可の記載が無い | 2. 入力 | 公開情報ではない |
| 4 | API-007 | WANT | dba | 3. 処理 1 | 参照テーブルと検索条件・使うインデックスが書かれていない | 3. 処理 1 | DB を使う API |
| 5 | NEW | WANT | ops | 3. 処理 2 | from の日付はどのタイムゾーンで判定するか（日付の境界） | 2. 入力 | 既存の観点に無い |
| 6 | API-005 | Q | infra | 全体 | 1 会員の注文件数の上限はどの程度を想定しているか | 記載なし | 非機能要件書の参照も無い |
"""

# office-review の data/history.jsonl と同じ形式（decision と check_id は必須。無い行は警告して飛ばす）
HISTORY = [
    {
        "key": "2026-10-02_order-api_result.md#1",
        "decision": "adopted",
        "check_id": "ALL-003",
        "severity": "MUST",
        "role": "spec_manager",
        "doc_type": "api",
        "finding": "「適宜ページングする」では 1 ページの件数と並び順の安定性が決まらない",
        "triaged_at": "2026-10-02",
    },
    {
        "key": "2026-10-02_order-api_result.md#4",
        "decision": "rejected",
        "check_id": "API-007",
        "severity": "WANT",
        "role": "dba",
        "doc_type": "api",
        "finding": "参照テーブルと検索条件・使うインデックスが書かれていない",
        "triaged_at": "2026-10-02",
    },
    {
        "key": "2026-10-02_order-api_result.md#5",
        "decision": "adopted",
        "check_id": "NEW",
        "severity": "WANT",
        "role": "ops",
        "doc_type": "api",
        "finding": "from の日付はどのタイムゾーンで判定するか（日付の境界）",
        "triaged_at": "2026-10-02",
    },
]


def generate(out: Path) -> list[Path]:
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "指摘一覧"
    header = ["No", "成果物", "工程", "指摘者", "作成者", "重要度", "指摘内容", "対応内容", "判定", "指摘日"]
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(1, c).font = Font(bold=True)
        ws.cell(1, c).fill = PatternFill("solid", fgColor="DDEBF7")
    for i, (d, ph, rv, au, sev, text, dec) in enumerate(ROWS, start=1):
        day = f"2026-{'05' if ph == '要件定義' else '08'}-{10 + i % 15:02d}"
        ws.append([i, d, ph, rv, au, sev, text, "修正しました" if dec == "対応する" else "現状のままとします", dec, day])
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["G"].width = 70
    xlsx = out / "指摘一覧.xlsx"
    wb.save(xlsx)

    pulls = out / "github_pulls.json"
    pulls.write_text(json.dumps(PULLS, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")

    rd = out / "office-review/reviews"
    rd.mkdir(parents=True, exist_ok=True)
    (rd / "2026-10-02_order-api_result.md").write_text(REVIEW_MD, encoding="utf-8", newline="\n")
    hist = out / "office-review/history.jsonl"
    hist.write_text("".join(json.dumps(h, ensure_ascii=False) + "\n" for h in HISTORY), encoding="utf-8", newline="\n")
    return [xlsx, pulls, rd / "2026-10-02_order-api_result.md", hist]


if __name__ == "__main__":
    for p in generate(HERE):
        print(p)
