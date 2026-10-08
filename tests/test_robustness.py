"""採点で見つかった不具合の再現テスト（id の伏せ字・数式注入・表の id・時刻・費用の上限）. 名前や社名はすべて架空."""

from __future__ import annotations

import contextlib
import io
import json
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from devtools_common.cli import run_main
from devtools_common.httpclient import HttpClient
from fakes import EchoLLM

from quality_analytics import analyze, mask, model, report, taxonomy
from quality_analytics.classify import llm
from quality_analytics.cli import main
from quality_analytics.sources import github, table

ROOT = Path(__file__).resolve().parent.parent
CONFIG = ROOT / "config"
TAX = taxonomy.load(CONFIG / "taxonomy.yaml")


def run(tmp_path: Path, *args: str, config_dir: Path = CONFIG) -> tuple[int, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(
            main,
            ["--data", str(tmp_path / "data/findings.csv"), "--config-dir", str(config_dir), "--prompt-dir", str(ROOT / "prompts"), *args],
        )
    return rc, out.getvalue() + err.getvalue()


def _config_with_words(tmp_path: Path, words: list[str]) -> Path:
    conf = tmp_path / "config"
    for p in CONFIG.rglob("*"):
        if p.is_file():
            dst = conf / p.relative_to(CONFIG)
            dst.parent.mkdir(parents=True, exist_ok=True)
            dst.write_bytes(p.read_bytes())
    (conf / "mask.yaml").write_text(
        f"pseudonymize: [reviewer, author]\nwords: {json.dumps(words, ensure_ascii=False)}\npatterns: []\n", encoding="utf-8"
    )
    (conf / "sources/table.yaml").write_text(
        "columns: {id: 'No', text: 指摘内容, phase: 工程}\nfixed: {source: table}\nphase_map: {基本設計: BD}\n", encoding="utf-8"
    )
    return conf


def collect(tmp_path: Path, *files: Path, words: list[str] | None = None) -> tuple[int, str]:
    conf = _config_with_words(tmp_path, words or [])
    return run(tmp_path, "collect", "table", *map(str, files), config_dir=conf)


def _csv(path: Path, rows: list[tuple[str, str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("No,指摘内容,工程\n" + "".join(f"{i},{t},基本設計\n" for i, t in rows), encoding="utf-8")
    return path


# --- 高: id の伏せ字で指摘が消える ---------------------------------------------------------------


def test_masked_ids_do_not_collide(tmp_path: Path):
    a = _csv(tmp_path / "in/指摘一覧_A社.csv", [("1", "桁が無い")])
    b = _csv(tmp_path / "in/指摘一覧_B社.csv", [("1", "型が無い")])
    rc, out = collect(tmp_path, a, b, words=["A社", "B社"])
    assert rc == 0, out
    df = model.load(tmp_path / "data/findings.csv")
    assert len(df) == 2 and df["id"].is_unique
    assert set(df["text"]) == {"桁が無い", "型が無い"}
    # id にも社名は残らない（伏せた部分は安定したハッシュで区別する）
    assert not any("A社" in i or "B社" in i for i in df["id"])
    # 同じ入力をもう一度取り込んでも id は変わらない（伏せ字は何度かけても同じ）
    rc, out = collect(tmp_path, a, b, words=["A社", "B社"])
    assert rc == 0 and "追加 0 件" in out
    assert list(model.load(tmp_path / "data/findings.csv")["id"]) == list(df["id"])


def test_mask_ident_is_stable_and_idempotent():
    m = mask.Masker(words=["A社"])
    once = m.ident("table:指摘一覧_A社:1")
    assert "A社" not in once and once.endswith(":1") and once.startswith("table:")
    assert m.ident(once) == once
    assert m.ident("table:指摘一覧_A社:1") == once
    assert m.ident("table:plain:1") == "table:plain:1"


def test_merge_stops_on_conflicting_duplicate_ids():
    new = model.normalize(pd.DataFrame([{"id": "x:1", "text": "一つ目"}, {"id": "x:1", "text": "二つ目"}]))
    with pytest.raises(model.DuplicateIdError, match="x:1"):
        model.merge(model.empty(), new)
    # 中身まで同じ行（同じファイルを 2 回指定など）は 1 件にまとめる
    same = model.normalize(pd.DataFrame([{"id": "x:1", "text": "同じ"}, {"id": "x:1", "text": "同じ"}]))
    merged, added, _, _ = model.merge(model.empty(), same)
    assert (len(merged), added) == (1, 1)


def test_duplicate_ids_in_table_are_usage_error(tmp_path: Path):
    p = _csv(tmp_path / "in/list.csv", [("1", "一つ目"), ("1", "二つ目")])
    rc, out = collect(tmp_path, p)
    assert rc == 2 and "id が重複" in out
    assert not (tmp_path / "data/findings.csv").exists()


# --- 中: 数式注入 ------------------------------------------------------------------------------

EVIL = ['=HYPERLINK("https://example.invalid","x")', "+1+1", "-1+1", "@SUM(A1)", "\t=1", "\r=1"]


def test_excel_writes_formula_like_text_as_string(tmp_path: Path):
    df = model.normalize(pd.DataFrame([{"id": f"t:{i}", "phase": "BD", "text": t, "reason": t} for i, t in enumerate(EVIL)]))
    a = analyze.run(df, TAX)
    path = tmp_path / "summary.xlsx"
    report.excel(a, df, path)
    ws = openpyxl.load_workbook(path)["指摘"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        for col in ("text", "reason"):
            cell = row[head.index(col)]
            assert cell.data_type == "s", cell.value
    assert {r[head.index("text")].value for r in ws.iter_rows(min_row=2)} == set(df["text"])  # 値は変えない


def test_findings_csv_is_neutralized_and_round_trips(tmp_path: Path):
    df = model.normalize(pd.DataFrame([{"id": f"t:{i}", "text": t} for i, t in enumerate([*EVIL, "'先頭が引用符", "ふつう"])]))
    path = tmp_path / "findings.csv"
    model.save(df, path)
    raw = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    assert all(not v.startswith(("=", "+", "-", "@", "\t", "\r")) for v in raw["text"])
    assert list(model.load(path)["text"]) == list(df["text"])


# --- 中: 表の id の振り直し・削除した行 -----------------------------------------------------------


def test_blank_id_rows_get_stable_ids_and_deleted_rows_are_removed(tmp_path: Path):
    p = tmp_path / "in/list.csv"
    _csv(p, [("1", "一つ目"), ("", "番号なし"), ("3", "三つ目")])
    assert collect(tmp_path, p)[0] == 0
    first = model.load(tmp_path / "data/findings.csv").set_index("text")["id"].to_dict()
    assert first["一つ目"] == "table:list:1" and first["三つ目"] == "table:list:3"
    # 行を削除・挿入して取り込み直す（番号のある行の id は変わらない。番号なしの行も同じ id）
    _csv(p, [("", "新しい行"), ("", "番号なし"), ("3", "三つ目")])
    rc, out = collect(tmp_path, p)
    assert rc == 0 and "削除 1 件" in out, out
    df = model.load(tmp_path / "data/findings.csv")
    assert sorted(df["text"]) == sorted(["新しい行", "番号なし", "三つ目"])  # 二重に数えない
    ids = df.set_index("text")["id"].to_dict()
    assert ids["番号なし"] == first["番号なし"] and ids["三つ目"] == "table:list:3"


def test_other_sources_are_not_deleted_on_table_reimport(tmp_path: Path):
    p = _csv(tmp_path / "in/list.csv", [("1", "一つ目")])
    q = _csv(tmp_path / "in/other.csv", [("1", "別の表")])
    assert collect(tmp_path, p, q)[0] == 0
    _csv(p, [("2", "差し替え")])
    assert collect(tmp_path, p)[0] == 0
    assert set(model.load(tmp_path / "data/findings.csv")["text"]) == {"差し替え", "別の表"}


def test_same_text_without_id_gets_distinct_ids(tmp_path: Path):
    p = _csv(tmp_path / "in/list.csv", [("", "同じ"), ("", "同じ")])
    df = table.read(p, {"columns": {"id": "No", "text": "指摘内容"}})
    assert df["id"].is_unique and len(df) == 2


# --- 低: GitHub の日付・ページ・再試行 ----------------------------------------------------------


def test_github_dates_use_local_timezone():
    data = {
        "repo": "o/r",
        "pulls": [{"number": 1, "user": {"login": "a"}, "comments": [{"id": 1, "body": "x", "created_at": "2026-10-04T23:30:00Z"}]}],
    }
    assert github.to_findings(data).iloc[0]["created_at"] == "2026-10-05"  # 既定は Asia/Tokyo
    assert github.to_findings(data, tz="UTC").iloc[0]["created_at"] == "2026-10-04"


def test_github_download_stops_paging_at_since():
    calls: list[str] = []

    def fetch(url: str):
        calls.append(url)
        if "/pulls?" in url:
            page = int(url.rsplit("page=", 1)[1])
            day = 30 - page  # 新しい順. ページが進むほど古い
            return [{"number": page * 100 + i, "updated_at": f"2026-09-{day:02d}T00:00:00Z"} for i in range(100)]
        return []

    data = github.download("o/r", fetch, since="2026-09-28", max_pulls=1000)
    assert {p["number"] // 100 for p in data["pulls"]} == {1, 2}
    assert sum("/pulls?" in c for c in calls) == 3  # 全 50 ページは取らない


def test_github_since_uses_local_date():
    def fetch(url: str):
        if "/pulls?" in url:
            return [{"number": 1, "updated_at": "2026-09-30T16:00:00Z"}] if url.endswith("page=1") else []
        return []

    assert len(github.download("o/r", fetch, since="2026-10-01")["pulls"]) == 1  # 日本時間では 10/1 01:00


def test_github_http_fetch_retries_with_httpclient():
    import urllib.error

    attempts: list[str] = []

    class _Resp(io.BytesIO):
        status = 200
        headers: dict = {}

    def opener(req, timeout):
        attempts.append(req.full_url)
        if len(attempts) == 1:
            raise urllib.error.HTTPError(req.full_url, 502, "bad gateway", {}, None)  # type: ignore[arg-type]
        assert req.get_header("Authorization") == "Bearer t"
        return _Resp(b"[]")

    client = HttpClient(github.API, headers={"Authorization": "Bearer t"}, opener=opener, sleep=lambda s: None)
    fetch = github.http_fetch("t", client=client)
    assert fetch(f"{github.API}/repos/o/r/pulls?page=1") == []
    assert len(attempts) == 2


# --- 低: 実際の費用で止める ---------------------------------------------------------------------


def _many(n: int) -> pd.DataFrame:
    return model.normalize(pd.DataFrame([{"id": f"t:{i}", "phase": "BD", "text": f"指摘 {i} の本文"} for i in range(n)]))


def test_llm_run_stops_when_actual_cost_reaches_limit(tmp_path: Path):
    conf = {"model": "claude-haiku-4-5-20251001", "batch_size": 2}
    p = llm.price(conf["model"], conf)
    assert p is not None
    client = EchoLLM(lambda item: {"type": "C1", "topic": "", "root_cause": "", "reason": ""}, model=conf["model"])
    one_call = (100 * p.input + 50 * p.output) / 1_000_000  # tool_reply の既定の usage
    out, res = llm.run(_many(6), TAX, conf, client, tmp_path / "cache.jsonl", ROOT / "prompts", max_cost=one_call * 1.5, price=p)
    assert res.calls == 1 and res.stopped and res.skipped == 4
    assert (out["type"] == "C1").sum() == 2
    assert res.cost(p) <= one_call * 1.5


def test_cli_reports_stop_by_actual_cost(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    p = _csv(tmp_path / "in/list.csv", [(str(i), f"指摘 {i}") for i in range(1, 46)])
    assert collect(tmp_path, p)[0] == 0
    monkeypatch.setattr(
        llm,
        "make_client",
        lambda conf: EchoLLM(lambda item: {"type": "C1", "topic": "", "root_cause": "", "reason": ""}, model=conf["model"]),
    )
    # 見積もりは上限内だが、実際の費用（偽の LLM の usage を大きくする）で途中で止まる
    from devtools_common.llm import Usage

    orig = llm.RunResult.add

    def add(self, usage):
        orig(self, Usage(input_tokens=2_000_000, output_tokens=0, cache_creation_input_tokens=0, cache_read_input_tokens=0))

    monkeypatch.setattr(llm.RunResult, "add", add)
    rc, out = run(tmp_path, "classify", "--mode", "api", "--yes", "--max-cost", "3")
    assert rc == 3 and "費用の上限" in out and "止めました" in out, out
    df = model.load(tmp_path / "data/findings.csv")
    assert 0 < (df["type"] == "C1").sum() < 45  # 分類できた分は保存する


def test_estimate_includes_retry_margin(tmp_path: Path):
    conf = {"model": "claude-haiku-4-5-20251001", "batch_size": 20}
    p = llm.price(conf["model"], conf)
    est = llm.estimate(_many(3), TAX, conf, ROOT / "prompts")
    assert est.cost_with_margin(p, conf) == pytest.approx(est.cost(p) * (1 + llm.DEFAULT_RETRY_MARGIN))
    assert est.cost_with_margin(p, {**conf, "retry_margin": 0}) == pytest.approx(est.cost(p))


# --- 第 2 回: キャッシュの壊れた行・GitHub の打ち切り・通信の失敗・ログの書式 ---------------------------------


def test_llm_cache_skips_broken_lines_and_warns(tmp_path: Path, caplog: pytest.LogCaptureFixture):
    """キャッシュの最終行が書きかけでも、以後の実行が失敗しない（壊れた行は読み飛ばして分類し直す）."""
    conf = {"model": "claude-haiku-4-5-20251001", "batch_size": 20}
    cache = tmp_path / "cache.jsonl"
    client = EchoLLM(lambda item: {"type": "C1", "topic": "", "root_cause": "", "reason": ""}, model=conf["model"])
    llm.run(_many(2), TAX, conf, client, cache, ROOT / "prompts")
    good = cache.read_text(encoding="utf-8").splitlines()
    assert len(good) == 2
    # 途中で止まった: 2 行目が書きかけ（改行も無い）・不正な行も混ざる
    cache.write_text(good[0] + "\n" + '["not", "a", "record"]\n' + good[1][: len(good[1]) // 2], encoding="utf-8")
    client = EchoLLM(lambda item: {"type": "C1", "topic": "", "root_cause": "", "reason": ""}, model=conf["model"])
    with caplog.at_level("WARNING", logger="quality_analytics.classify.llm"):
        out, res = llm.run(_many(2), TAX, conf, client, cache, ROOT / "prompts")
    assert res.cached == 1 and res.classified == 1 and (out["type"] == "C1").all()
    assert "壊れた行 2 行を読み飛ばしました" in caplog.text
    # 書きかけの行の後ろに追記しても、新しい行は壊れない（改行を足してから追記する）
    assert len(llm._load_cache(cache)) == 2


def test_github_download_warns_when_max_pulls_cuts(caplog: pytest.LogCaptureFixture):
    def fetch(url: str):
        if "/pulls?" in url:
            page = int(url.rsplit("page=", 1)[1])
            return [{"number": page * 100 + i, "updated_at": "2026-09-30T00:00:00Z"} for i in range(100)] if page <= 3 else []
        return []

    with caplog.at_level("WARNING", logger="quality_analytics.sources.github"):
        data = github.download("o/r", fetch, max_pulls=150)
    assert len(data["pulls"]) == 150
    assert any("150 件で打ち切りました" in n for n in data["truncated"])
    assert "PR を 150 件で打ち切りました" in caplog.text


def test_github_download_warns_when_pages_run_out(caplog: pytest.LogCaptureFixture):
    """コメントが 50 ページ（5000 件）を超える PR: 黙って捨てずに警告し、件数を残す."""

    def fetch(url: str):
        if "/pulls?" in url:
            return [{"number": 7, "updated_at": "2026-09-30T00:00:00Z"}] if url.endswith("page=1") else []
        if "/comments" in url:
            return [{"id": i, "body": "x"} for i in range(github.PER_PAGE)]  # いつまでも満杯
        return []

    with caplog.at_level("WARNING", logger="quality_analytics.sources.github"):
        data = github.download("o/r", fetch)
    assert len(data["pulls"][0]["comments"]) == github.MAX_PAGES * github.PER_PAGE
    assert data["truncated"] == [
        f"pulls/7/comments は {github.MAX_PAGES} ページ（{github.MAX_PAGES * github.PER_PAGE} 件）で打ち切りました"
    ]
    assert "取得を打ち切りました" in caplog.text


def test_github_download_without_truncation_has_no_notes():
    def fetch(url: str):
        if "/pulls?" in url:
            return [{"number": 1, "updated_at": "2026-09-30T00:00:00Z"}] if url.endswith("page=1") else []
        return []

    assert "truncated" not in github.download("o/r", fetch)


def test_cli_github_prints_truncation(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("GH_TOKEN", "dummy")
    monkeypatch.setattr(github, "http_fetch", lambda token: object())
    monkeypatch.setattr(
        github, "download", lambda repo, fetch, since="", tz="": {"repo": repo, "pulls": [], "truncated": ["PR は 200 件で打ち切りました"]}
    )
    rc, out = run(tmp_path, "collect", "github", "--repo", "o/r")
    assert rc == 0 and "注意: PR は 200 件で打ち切りました" in out and "PR 0 件を取得しました" in out


class _FlakyLLM(EchoLLM):
    """最初の ok 回だけ応答し、その後は接続できない（再試行しても失敗する）."""

    def __init__(self, ok: int, **kw):
        super().__init__(lambda item: {"type": "C1", "topic": "", "root_cause": "", "reason": ""}, max_retries=0, **kw)
        self._ok = ok

    def _open(self, req, timeout):
        import urllib.error

        if self._ok <= 0:
            raise urllib.error.URLError("connection reset")
        self._ok -= 1
        return super()._open(req, timeout)


def test_llm_run_stops_on_connection_error_and_keeps_results(tmp_path: Path):
    conf = {"model": "claude-haiku-4-5-20251001", "batch_size": 2}
    out, res = llm.run(_many(6), TAX, conf, _FlakyLLM(1, model=conf["model"]), tmp_path / "cache.jsonl", ROOT / "prompts")
    assert res.error.startswith("APIConnectionError") and res.skipped == 4
    assert (out["type"] == "C1").sum() == 2 and res.classified == 2


def test_cli_saves_partial_results_on_connection_error(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    p = _csv(tmp_path / "in/list.csv", [(str(i), f"指摘 {i}") for i in range(1, 46)])
    assert collect(tmp_path, p)[0] == 0
    monkeypatch.setenv("ANTHROPIC_API_KEY", "dummy")
    monkeypatch.setattr(llm, "make_client", lambda conf: _FlakyLLM(1, model=conf["model"]))
    rc, out = run(tmp_path, "classify", "--mode", "api", "--yes")
    assert rc == 3 and "LLM API の呼び出しに失敗" in out and "保存しました" in out, out
    df = model.load(tmp_path / "data/findings.csv")
    assert 0 < (df["type"] == "C1").sum() < 45  # 通信の失敗の前に分類できた分は findings.csv に残る


def test_numeric_log_format_survives_redact_filter(tmp_path: Path):
    """devtools-common の RedactFilter で %d のログが消えていた件（28bab4c で修正）の再現テスト."""
    import logging

    from devtools_common.log import RedactFilter

    buf = io.StringIO()
    handler = logging.StreamHandler(buf)
    handler.addFilter(RedactFilter())
    logger = logging.getLogger("quality_analytics.classify.llm")
    logger.addHandler(handler)
    old = logger.level
    logger.setLevel(logging.INFO)
    err = io.StringIO()
    try:
        with contextlib.redirect_stderr(err):
            conf = {"model": "claude-haiku-4-5-20251001", "batch_size": 20}
            client = EchoLLM(lambda item: {"type": "C1", "topic": "", "root_cause": "", "reason": ""}, model=conf["model"])
            llm.run(_many(1), TAX, conf, client, tmp_path / "cache.jsonl", ROOT / "prompts")
    finally:
        logger.removeHandler(handler)
        logger.setLevel(old)
    assert "llm usage: calls=1 in=100 out=50" in buf.getvalue()
    assert "Logging error" not in err.getvalue()
